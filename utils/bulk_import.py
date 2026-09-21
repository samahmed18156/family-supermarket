"""
FREE Bulk Import - Excel/CSV to products.json
No paid services, uses openpyxl (free)

Usage:
  pip install openpyxl
  python utils/bulk_import.py your_products.xlsx
  or
  python utils/bulk_import.py your_products.csv

Excel format:
  | id | name | category | price | unit | image | in_stock | special | special_price |
"""
import sys, json, csv
from pathlib import Path


def import_excel(path):
    try:
        import openpyxl
    except ImportError:
        print("❌ Need openpyxl: pip install openpyxl")
        return

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("Empty file")
        return

    headers = [str(h).lower().strip() if h else "" for h in rows[0]]
    products = []

    for row in rows[1:]:
        if not any(row):
            continue
        data = dict(zip(headers, row))
        try:
            product = {
                "id": int(data.get("id", len(products) + 1)),
                "name": str(data.get("name", "")).strip(),
                "category": str(data.get("category", "General")).strip(),
                "price": float(data.get("price", 0)),
                "unit": str(data.get("unit", "unit")).strip(),
                "image": str(data.get("image", "rice.jpg")).strip(),
                "in_stock": str(data.get("in_stock", "true")).lower() in ("true", "1", "yes"),
                "special": str(data.get("special", "false")).lower() in ("true", "1", "yes"),
            }
            if product["special"] and data.get("special_price"):
                product["special_price"] = float(data.get("special_price"))
            if product["name"]:
                products.append(product)
        except Exception as e:
            print(f"Skip row {data}: {e}")

    Path("../products.json").write_text(json.dumps(products, indent=2))
    print(f"✅ Imported {len(products)} products to products.json")


def import_csv(path):
    products = []
    with open(path, newline='', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for i, row in enumerate(reader, 1):
            try:
                product = {
                    "id": int(row.get("id", i)),
                    "name": row.get("name", "").strip(),
                    "category": row.get("category", "General").strip(),
                    "price": float(row.get("price", 0)),
                    "unit": row.get("unit", "unit").strip(),
                    "image": row.get("image", "rice.jpg").strip(),
                    "in_stock": row.get("in_stock", "true").lower() in ("true", "1", "yes"),
                    "special": row.get("special", "false").lower() in ("true", "1", "yes"),
                }
                if product["special"] and row.get("special_price"):
                    product["special_price"] = float(row["special_price"])
                if product["name"]:
                    products.append(product)
            except Exception as e:
                print(f"Skip row {row}: {e}")

    Path("../products.json").write_text(json.dumps(products, indent=2))
    print(f"✅ Imported {len(products)} products to products.json")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python utils/bulk_import.py file.xlsx or file.csv")
        print("\nCreating example template...")
        try:
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["id", "name", "category", "price", "unit", "image", "in_stock", "special", "special_price"])
            ws.append([1, "Brown Bread", "Bakery", 18.5, "loaf", "bread.jpg", True, False, ""])
            ws.append([2, "Milk 2L", "Dairy", 32, "bottle", "milk.jpg", True, True, 27])
            ws.append([3, "Rice 10kg", "Staples", 129.99, "bag", "rice.jpg", True, False, ""])
            wb.save("products_template.xlsx")
            print(
                "✅ Created products_template.xlsx - edit it and run: python utils/bulk_import.py products_template.xlsx")
        except:
            print("Install openpyxl to create template: pip install openpyxl")
        sys.exit(0)

    path = Path(sys.argv[1])
    if path.suffix.lower() in [".xlsx", ".xls"]:
        import_excel(path)
    elif path.suffix.lower() == ".csv":
        import_csv(path)
    else:
        print("Unsupported file. Use .xlsx or .csv")
